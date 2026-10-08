import json
import os
from collections import Counter
from datetime import datetime, timezone

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW_FILE = os.path.join(BASE_DIR, "raw", "hulpdienstvoertuigenbenelux_raw.json")
DISCORD_WEBHOOK_URL = os.getenv("DISCORD_WEBHOOK_URL")  # Loaded from GitHub Secrets

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def fetch_and_check():
    print(f"Loading raw data from {RAW_FILE}...")

    try:
        with open(RAW_FILE, encoding="utf-8") as f:
            records = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        print(f"Failed to load raw data: {e}")
        return

    headers_row = []
    for record in records:
        for key in record:
            if key not in headers_row:
                headers_row.append(key)

    onbekend_entries = [
        record for record in records
        if any("ONBEKEND" in str(value).upper() for value in record.values())
    ]

    print(f"Found {len(onbekend_entries)} rows containing 'ONBEKEND'.")

    if onbekend_entries:
        # Sla op als Excel-bestand: blad 1 overzicht, blad 2 de rijen zelf
        filename = f"onbekend_{datetime.now(timezone.utc):%Y-%m-%d}.xlsx"
        save_to_xlsx(onbekend_entries, headers_row, filename)

        # Verstuur naar Discord als bestand
        send_discord_alert_with_file(onbekend_entries, filename)
    else:
        print("No 'ONBEKEND' values found!")


def _autosize(sheet):
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(width + 2, 60)


def _regio_sort_key(regio):
    if regio == "(geen regio)":
        return (2, 0, regio)
    nummer, _, naam = regio.partition(" - ")
    if nummer.isdigit():
        return (0, int(nummer), naam)
    return (1, 0, regio)


def save_to_xlsx(entries, headers, filename):
    print(f"Saving results to {filename}...")
    workbook = Workbook()

    # Blad 1: overzicht per regio, met per veld hoe vaak ONBEKEND voorkomt
    per_veld = Counter(
        key for entry in entries for key, value in entry.items()
        if "ONBEKEND" in str(value).upper()
    )
    velden = [veld for veld, _ in per_veld.most_common()]

    per_regio = {}
    for entry in entries:
        regio = entry.get("Regio", "") or "(geen regio)"
        counts = per_regio.setdefault(regio, Counter())
        counts["Totaal"] += 1
        for key, value in entry.items():
            if "ONBEKEND" in str(value).upper():
                counts[key] += 1

    overzicht = workbook.active
    overzicht.title = "Overzicht"
    overzicht.append(["Regio", "Totaal rijen"] + velden)
    for cell in overzicht[1]:
        cell.font = Font(bold=True)
    for regio in sorted(per_regio, key=_regio_sort_key):
        counts = per_regio[regio]
        overzicht.append([regio, counts["Totaal"]] + [counts[veld] or None for veld in velden])
    overzicht.auto_filter.ref = overzicht.dimensions
    overzicht.append(["Totaal", len(entries)] + [per_veld[veld] for veld in velden])
    for cell in overzicht[overzicht.max_row]:
        cell.font = Font(bold=True)
    overzicht.freeze_panes = "B2"
    _autosize(overzicht)

    # Blad 2: alle rijen met ONBEKEND. Alles als tekst, zodat Excel
    # roepnummers zoals '01-81' niet als datum (Jan-81) leest.
    data = workbook.create_sheet("Onbekend")
    data.append(headers)
    for cell in data[1]:
        cell.font = Font(bold=True)
    for entry in entries:
        data.append([str(entry.get(column, "") or "") for column in headers])
    for row in data.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"
    data.freeze_panes = "A2"
    data.auto_filter.ref = data.dimensions
    _autosize(data)

    workbook.save(filename)


def send_discord_alert_with_file(entries, filename):
    if not DISCORD_WEBHOOK_URL:
        print(f"Discord Webhook URL not set. File is saved locally as {filename}.")
        return

    print("Sending file to Discord...")

    payload = {
        "content": (
            "🔎 **Wekelijkse controle: ontbrekende gegevens**\n"
            f"Bij **{len(entries)} rijen** staat nog `ONBEKEND` ingevuld. "
            "In de bijlage vind je een overzicht en de volledige lijst."
        )
    }

    try:
        with open(filename, "rb") as f:
            files = {
                "file": (filename, f, XLSX_MIME)
            }
            response = requests.post(DISCORD_WEBHOOK_URL, data=payload, files=files)
            response.raise_for_status()
            print("Successfully sent alert and file to Discord.")
    except Exception as e:
        print(f"Failed to send Discord alert with file: {e}")


if __name__ == "__main__":
    fetch_and_check()
